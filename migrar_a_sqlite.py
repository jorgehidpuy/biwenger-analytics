# -*- coding: utf-8 -*-
"""Pasa los datos del Excel a la base SQLite biwenger.db (se ejecuta UNA vez).

    python migrar_a_sqlite.py [ruta del Excel]

No toca el Excel. Si biwenger.db ya existe, no hace nada salvo que se le pase
--forzar (y entonces la deja aparte como biwenger.db.anterior).

Qué se lleva y qué no:
  - Las hojas de datos, con sus columnas tal cual. Las celdas que eran FÓRMULAS
    se dejan vacías: ahora esos valores los calcula calculos.py.
  - ClasificacionCalendario: solo el calendario (columnas A-J); la
    clasificación de al lado era todo fórmulas.
  - Historico: de bloques de columnas por día a una tabla de una fila por
    jugador y día.
  - NO se llevan Mercado_Global, Calcs y Scouting_Helper (solo fórmulas) ni
    Equipos (listas para los desplegables del Excel).
"""
import os
import shutil
import sys
from datetime import datetime

import openpyxl

import almacen

HOJAS_FUERA = {"Mercado_Global", "Calcs", "Scouting_Helper", "Equipos", "Historico"}
FILA_CABECERA = {"ClasificacionCalendario": 2}
COLUMNAS_MAX = {"ClasificacionCalendario": 10}


def _es_formula(valor):
    """Fórmulas normales ("=SUM(...)") y matriciales (objetos ArrayFormula)."""
    if isinstance(valor, str):
        return valor.startswith("=")
    return type(valor).__name__ in ("ArrayFormula", "DataTableFormula")


def migrar(ruta_excel, ruta_db=None, forzar=False, registro=print):
    ruta_db = ruta_db or almacen.RUTA_DB
    if os.path.exists(ruta_db):
        if not forzar:
            registro(f"{ruta_db} ya existe: no se toca. Usa --forzar para rehacerla.")
            return False
        shutil.move(ruta_db, ruta_db + ".anterior")
        registro(f"La base anterior queda como {ruta_db}.anterior")

    wb = openpyxl.load_workbook(ruta_excel)
    con = almacen.conectar(ruta_db)
    resumen = []
    try:
        with con:
            for ws in wb.worksheets:
                if ws.title in HOJAS_FUERA:
                    continue
                max_col = min(ws.max_column, COLUMNAS_MAX.get(ws.title, ws.max_column))
                rejilla = [[None if _es_formula(v) else v for v in fila]
                           for fila in ws.iter_rows(max_col=max_col, values_only=True)]
                hoja = almacen.Hoja(None, ws.title, rejilla, FILA_CABECERA.get(ws.title, 1))
                rotulos, datos = hoja.a_tabla()
                almacen._escribir_tabla(con, ws.title, rotulos, datos, hoja.fila_cabecera)
                resumen.append(f"{ws.title}: {len(datos)} filas x {len(rotulos)} columnas")

        dias = 0
        if "Historico" in wb.sheetnames:
            dias = _migrar_historico(wb["Historico"], con)
            resumen.append(f"Historico: {dias} días")
    finally:
        con.close()
    for linea in resumen:
        registro("   " + linea)
    return True


def _migrar_historico(ws, con):
    """Mismo recorrido que _leer_bloques_historico de la app: bloques de 7
    columnas desde la C, fecha en la fila 2, rótulos en la 3, datos desde la 4."""
    filas = [list(f) for f in ws.iter_rows(values_only=True)]

    def celda(r, c):
        return filas[r - 1][c - 1] if r <= len(filas) and c <= len(filas[r - 1]) else None

    ancho = len(almacen.HIST_COLS)
    max_col = max((len(f) for f in filas), default=0)
    col, dias = 3, 0
    while col <= max_col:
        if celda(3, col) in (None, ""):
            col += 1
            continue
        fecha = celda(2, col)
        if not isinstance(fecha, datetime):
            fecha = datetime.strptime(str(fecha).strip(), "%d/%m/%Y")
        bloque = []
        r = 4
        while celda(r, col) not in (None, ""):
            bloque.append({campo: celda(r, col + i) for i, campo in enumerate(almacen.HIST_COLS)})
            r += 1
        if almacen.historico_anadir_dia(fecha, bloque, con=con):
            dias += 1
        col += ancho + 1
    return dias


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    ruta = args[0] if args else os.path.join(almacen.CARPETA, "Biwenger_V19_.xlsx")
    print(f"Migrando {ruta} -> {almacen.RUTA_DB}")
    migrar(ruta, forzar="--forzar" in sys.argv)
